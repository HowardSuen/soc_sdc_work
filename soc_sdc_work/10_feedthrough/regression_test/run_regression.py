#!/usr/bin/env python3
"""Complex regression for 10_extract_feedthrough.py."""

from __future__ import print_function

import csv
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook


BASE = Path(__file__).resolve().parent
SOC = BASE.parent.parent
EX10 = SOC / "10_feedthrough" / "10_extract_feedthrough.py"
WORK = BASE / "work_complex"

REQ = [
    "Input",
    "Input Width",
    "Input Used Width",
    "From Whom",
    "Output",
    "Output Width",
    "Output Used Width",
    "To Top",
    "Inout",
    "Inout Width",
    "Inout Connectivity",
    "Inout Name",
]


def clean_dir(path):
    if path.exists():
        shutil.rmtree(str(path))
    path.mkdir(parents=True)


def sh(args, cwd):
    return subprocess.run(
        [sys.executable] + [str(arg) for arg in args],
        cwd=str(cwd),
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        universal_newlines=True,
    )


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def write_info_all(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "info"
    ws.append(["inst_name", "module_name", "owner"])
    for row in rows:
        ws.append(row)
    wb.save(str(path))


def append_port_row(ws, **values):
    ws.append([values.get(col, "") for col in REQ])


def write_ports(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "u_a"
    ws.append(REQ)
    append_port_row(ws, Input="resp_valid", **{"Input Width": 1}, Output="req_data", **{"Output Width": 2})

    ws = wb.create_sheet("u_b")
    ws.append(REQ)
    append_port_row(
        ws,
        Input="fti_0_a2d_req_data",
        **{"Input Width": 2},
        Output="fto_0_a2d_req_data",
        **{"Output Width": 2}
    )
    append_port_row(
        ws,
        Input="fti_1_d2a_resp_valid",
        **{"Input Width": 1},
        Output="fto_1_d2a_resp_valid",
        **{"Output Width": 1}
    )
    append_port_row(ws, Input="cfg_i", **{"Input Width": 1}, Output="keep_o", **{"Output Width": 1})

    ws = wb.create_sheet("u_c")
    ws.append(REQ)
    append_port_row(
        ws,
        Input="fti_1_a2d_req_data",
        **{"Input Width": 2},
        Output="fto_1_a2d_req_data",
        **{"Output Width": 2}
    )
    append_port_row(
        ws,
        Input="fti_0_d2a_resp_valid",
        **{"Input Width": 1},
        Output="fto_0_d2a_resp_valid",
        **{"Output Width": 1}
    )
    append_port_row(ws, Input="spare_i", **{"Input Width": 1})

    ws = wb.create_sheet("u_d")
    ws.append(REQ)
    append_port_row(ws, Input="req_data", **{"Input Width": 2}, Output="resp_valid", **{"Output Width": 1})
    wb.save(str(path))


def write_connection_inventory(path):
    path.parent.mkdir(parents=True)
    headers = [
        "connection_id",
        "connection_type",
        "src_instance",
        "src_direction",
        "src_port",
        "src_bit_index",
        "src_endpoint_key",
        "src_soc_object",
        "dst_instance",
        "dst_direction",
        "dst_port",
        "dst_bit_index",
        "dst_endpoint_key",
        "dst_soc_object",
        "validation_status",
        "note",
    ]
    with path.open("w", encoding="utf-8", newline="") as file_obj:
        writer = csv.DictWriter(file_obj, fieldnames=headers)
        writer.writeheader()

        def add(src_i, src_d, src_p, dst_i, dst_d, dst_p, conn_type="harden_to_harden"):
            src_bit = src_p.split("[")[-1].rstrip("]") if "[" in src_p else ""
            dst_bit = dst_p.split("[")[-1].rstrip("]") if "[" in dst_p else ""
            cid = "CONN_%s_%s__%s_%s" % (
                src_i,
                src_p.replace("[", "_bit").replace("]", ""),
                dst_i,
                dst_p.replace("[", "_bit").replace("]", ""),
            )
            writer.writerow({
                "connection_id": cid,
                "connection_type": conn_type,
                "src_instance": src_i,
                "src_direction": src_d,
                "src_port": src_p,
                "src_bit_index": src_bit,
                "src_endpoint_key": "%s:%s:%s" % (src_i, src_d, src_p),
                "src_soc_object": "%s/%s" % (src_i, src_p),
                "dst_instance": dst_i,
                "dst_direction": dst_d,
                "dst_port": dst_p,
                "dst_bit_index": dst_bit,
                "dst_endpoint_key": "%s:%s:%s" % (dst_i, dst_d, dst_p),
                "dst_soc_object": "%s/%s" % (dst_i, dst_p),
                "validation_status": "matched",
                "note": "",
            })

        for bit in ("0", "1"):
            add("u_a", "output", "req_data[%s]" % bit, "u_b", "input", "fti_0_a2d_req_data[%s]" % bit)
            add("u_b", "output", "fto_0_a2d_req_data[%s]" % bit, "u_c", "input", "fti_1_a2d_req_data[%s]" % bit)
            add("u_c", "output", "fto_1_a2d_req_data[%s]" % bit, "u_d", "input", "req_data[%s]" % bit)
        add("u_d", "output", "resp_valid", "u_c", "input", "fti_0_d2a_resp_valid")
        add("u_c", "output", "fto_0_d2a_resp_valid", "u_b", "input", "fti_1_d2a_resp_valid")
        add("u_b", "output", "fto_1_d2a_resp_valid", "u_a", "input", "resp_valid")


def write_pending(root):
    pending = root / "00_harden_port_inventory" / "pending"
    pending.mkdir(parents=True)
    (pending / "u_a.ports").write_text(
        "input resp_valid\n"
        "output req_data[0]\n"
        "output req_data[1]\n",
        encoding="utf-8",
    )
    (pending / "u_b.ports").write_text(
        "input fti_0_a2d_req_data[0]\n"
        "input fti_0_a2d_req_data[1]\n"
        "input fti_1_d2a_resp_valid\n"
        "input cfg_i\n"
        "output fto_0_a2d_req_data[0]\n"
        "output fto_0_a2d_req_data[1]\n"
        "output fto_1_d2a_resp_valid\n"
        "output keep_o\n",
        encoding="utf-8",
    )
    (pending / "u_c.ports").write_text(
        "input fti_1_a2d_req_data[0]\n"
        "input fti_1_a2d_req_data[1]\n"
        "input fti_0_d2a_resp_valid\n"
        "input spare_i\n"
        "output fto_1_a2d_req_data[0]\n"
        "output fto_1_a2d_req_data[1]\n"
        "output fto_0_d2a_resp_valid\n",
        encoding="utf-8",
    )
    (pending / "u_d.ports").write_text(
        "input req_data[0]\n"
        "input req_data[1]\n"
        "output resp_valid\n",
        encoding="utf-8",
    )


def build_positive_case(d):
    clean_dir(d)
    write_info_all(d / "info_all.xlsx", [
        ["u_a", "harden_a", "alice"],
        ["u_b", "harden_b", "bob"],
        ["u_c", "harden_c", "carol"],
        ["u_d", "harden_d", "dave"],
    ])
    write_ports(d / "ports_feedthrough.xlsx")
    write_connection_inventory(d / "00_harden_port_inventory" / "connection_inventory.csv")
    write_pending(d)


def run_positive_case():
    d = WORK / "positive"
    build_positive_case(d)

    result = sh([EX10], d)
    require(result.returncode == 0, "10 positive case failed:\n%s\n%s" % (result.stdout, result.stderr))

    inv = (d / "feedthrough_inventory.csv").read_text(encoding="utf-8")
    sdc = (d / "common" / "10_feedthrough.sdc").read_text(encoding="utf-8")
    report = (d / "feedthrough_check_report_common.txt").read_text(encoding="utf-8")
    pending_b = (d / "00_harden_port_inventory" / "pending" / "u_b.ports").read_text(encoding="utf-8")
    pending_c = (d / "00_harden_port_inventory" / "pending" / "u_c.ports").read_text(encoding="utf-8")
    removed = (d / "00_harden_port_inventory" / "removed_log" / "10_feedthrough.removed").read_text(encoding="utf-8")

    require("FT_u_b_0_a2d_req_data_bit1" in inv, "bit-level u_b req feedthrough id missing")
    require("FT_u_c_1_a2d_req_data_bit0" in inv, "bit-level u_c req feedthrough id missing")
    require("FT_u_c_0_d2a_resp_valid" in inv, "reverse scalar feedthrough id missing")
    require("segments with validation_status=matched: 6" in report, "matched segment coverage missing")
    require("removed 12 feedthrough port endpoint(s) from pending" in report, "pending removal summary missing")
    require("input fti_0_a2d_req_data[0]" not in pending_b, "u_b fti bit was not removed")
    require("output fto_1_d2a_resp_valid" not in pending_b, "u_b fto scalar was not removed")
    require("input cfg_i" in pending_b and "output keep_o" in pending_b, "non-feedthrough u_b ports should stay pending")
    require("input spare_i" in pending_c, "non-feedthrough u_c port should stay pending")
    require("u_b input fti_0_a2d_req_data[1] covered_by=10_feedthrough" in removed, "removed log missing u_b fti bit")
    require("feedthrough_id=FT_u_c_0_d2a_resp_valid" in removed, "removed log missing reverse scalar id")
    require("# FT_u_c_1_a2d_req_data_bit0 status=matched" in sdc, "SDC manifest missing bit segment")
    require("#   upstream [get_pins {u_b/fto_0_a2d_req_data[0]}]" in sdc, "SDC manifest missing upstream anchor")

    rerun = sh([EX10], d)
    require(rerun.returncode == 0, "10 rerun should be idempotent:\n%s\n%s" % (rerun.stdout, rerun.stderr))
    rerun_report = (d / "feedthrough_check_report_common.txt").read_text(encoding="utf-8")
    require("no previous_removed record exists" not in rerun_report, "10 pending update is not idempotent")


def run_unpaired_error_case():
    d = WORK / "unpaired"
    clean_dir(d)
    write_info_all(d / "info_all.xlsx", [["u_x", "harden_x", "x"]])
    wb = Workbook()
    ws = wb.active
    ws.title = "u_x"
    ws.append(REQ)
    append_port_row(ws, Input="fti_x2y_ctrl", **{"Input Width": 1})
    wb.save(str(d / "ports_x.xlsx"))
    result = sh([EX10, "--no-update-pending"], d)
    require(result.returncode == 1, "unpaired fti should fail")
    report = (d / "feedthrough_check_report_common.txt").read_text(encoding="utf-8")
    require("has no matching fto_x2y_ctrl" in report, "unpaired fti error missing")


def run_wrong_direction_error_case():
    d = WORK / "wrong_direction"
    clean_dir(d)
    write_info_all(d / "info_all.xlsx", [["u_x", "harden_x", "x"]])
    wb = Workbook()
    ws = wb.active
    ws.title = "u_x"
    ws.append(REQ)
    append_port_row(ws, Output="fti_x2y_ctrl", **{"Output Width": 1})
    append_port_row(ws, Input="fto_x2y_ctrl", **{"Input Width": 1})
    wb.save(str(d / "ports_x.xlsx"))
    result = sh([EX10, "--no-update-pending"], d)
    require(result.returncode == 1, "wrong fti/fto directions should fail")
    report = (d / "feedthrough_check_report_common.txt").read_text(encoding="utf-8")
    require("fti_x2y_ctrl is listed as output" in report, "wrong fti direction error missing")
    require("fto_x2y_ctrl is listed as input" in report, "wrong fto direction error missing")


def run_inout_direction_error_case():
    d = WORK / "inout_direction"
    clean_dir(d)
    write_info_all(d / "info_all.xlsx", [["u_x", "harden_x", "x"]])
    wb = Workbook()
    ws = wb.active
    ws.title = "u_x"
    ws.append(REQ)
    append_port_row(ws, Inout="fti_x2y_gpio", **{"Inout Width": 1})
    wb.save(str(d / "ports_x.xlsx"))
    result = sh([EX10, "--no-update-pending"], d)
    require(result.returncode == 1, "inout fti should fail in 10 v1")
    report = (d / "feedthrough_check_report_common.txt").read_text(encoding="utf-8")
    require("inout feedthrough ports must be split" in report, "inout direction explanation missing")


def run_index_gap_error_case():
    d = WORK / "index_gap"
    clean_dir(d)
    write_info_all(d / "info_all.xlsx", [
        ["u_b", "harden_b", "b"],
        ["u_c", "harden_c", "c"],
    ])
    wb = Workbook()
    ws = wb.active
    ws.title = "u_b"
    ws.append(REQ)
    append_port_row(ws, Input="fti_0_a2d_req", **{"Input Width": 1}, Output="fto_0_a2d_req", **{"Output Width": 1})
    ws = wb.create_sheet("u_c")
    ws.append(REQ)
    append_port_row(ws, Input="fti_2_a2d_req", **{"Input Width": 1}, Output="fto_2_a2d_req", **{"Output Width": 1})
    wb.save(str(d / "ports_gap.xlsx"))
    result = sh([EX10, "--no-update-pending"], d)
    require(result.returncode == 1, "index gap should fail")
    report = (d / "feedthrough_check_report_common.txt").read_text(encoding="utf-8")
    require("feedthrough indexes are not contiguous from 0" in report, "index gap error missing")


def run_scenario_error_case():
    d = WORK / "scenario"
    build_positive_case(d)
    result = sh([EX10, "-scenario", "func"], d)
    require(result.returncode == 1, "scenario-specific 10 feedthrough should fail in v1")
    report = (d / "feedthrough_check_report_func.txt").read_text(encoding="utf-8")
    require("supports only scenario=common" in report, "scenario common-only error missing")
    require(not (d / "scenarios" / "func_feedthrough.sdc").exists(), "scenario-specific feedthrough SDC should not be generated")


def main():
    clean_dir(WORK)
    run_positive_case()
    run_unpaired_error_case()
    run_wrong_direction_error_case()
    run_inout_direction_error_case()
    run_index_gap_error_case()
    run_scenario_error_case()
    print("10 feedthrough complex regression: PASS")
    print("  cases: positive, unpaired, wrong_direction, inout_direction, index_gap, scenario")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
